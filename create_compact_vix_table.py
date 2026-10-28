#!/usr/bin/env python3
"""Create the compact VIX regression table used in the paper body."""

from __future__ import annotations

from pathlib import Path

ROOT = Path(__file__).resolve().parent

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT = (
    ROOT
    / "answer"
    / "tables"
    / "Table_2_Association_Between_VIX_and_90-Day_Option-Implied_Measures.docx"
)


def set_cell_border(cell, **edges):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge_name, attributes in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in attributes.items():
            edge.set(qn(f"w:{key}"), str(value))


def write_cell(cell, text: str, *, bold: bool = False, centered: bool = False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
        if centered
        else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build_table():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "Table 2. Association Between VIX and 90-Day Option-Implied Measures"
    )
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    rows = [
        ("Dependent Variable", "SPX", "NDX"),
        ("ATM IV", "0.0557***", "0.0494***"),
        ("RR25", "−0.0263***", "−0.0246***"),
        ("Downside Tail Probability", "0.0286***", "0.0262***"),
        ("R² – ATM model", "0.944", "0.866"),
        ("R² – RR25 model", "0.823", "0.802"),
        ("R² – Tail model", "0.955", "0.874"),
    ]
    table = document.add_table(rows=len(rows), cols=3)
    table.autofit = False
    widths = [Inches(3.6), Inches(1.25), Inches(1.25)]

    for row_index, row_values in enumerate(rows):
        for column_index, value in enumerate(row_values):
            cell = table.cell(row_index, column_index)
            cell.width = widths[column_index]
            write_cell(
                cell,
                value,
                bold=row_index == 0,
                centered=column_index > 0 or row_index == 0,
            )
            no_border = {"val": "nil"}
            borders = {
                "top": no_border,
                "bottom": no_border,
                "left": no_border,
                "right": no_border,
                "insideH": no_border,
                "insideV": no_border,
            }
            if row_index == 0:
                borders["top"] = {"val": "single", "sz": "8", "color": "000000"}
                borders["bottom"] = {
                    "val": "single",
                    "sz": "6",
                    "color": "000000",
                }
            if row_index == len(rows) - 1:
                borders["bottom"] = {
                    "val": "single",
                    "sz": "8",
                    "color": "000000",
                }
            set_cell_border(cell, **borders)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.left_indent = Inches(0.05)
    note.paragraph_format.right_indent = Inches(0.05)
    note_run = note.add_run(
        "Notes: Each regression additionally controls for 21-day realized "
        "volatility, monthly index return, and median relative bid–ask spread. "
        "Predictors are standardized. HC1 heteroskedasticity-robust standard "
        "errors are used. N = 83 for each regression. *** p < 0.01."
    )
    note_run.font.name = "Times New Roman"
    note_run.font.size = Pt(9)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_table()
