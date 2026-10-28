#!/usr/bin/env python3
"""Create a Word report for the independent-IV robustness exercise."""

from __future__ import annotations

from pathlib import Path

ROOT = Path(__file__).resolve().parent

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_ROOT = ROOT / "outputs" / "iv_robustness_2018_2024"
CSV_ROOT = OUTPUT_ROOT / "csv"
CHART_ROOT = OUTPUT_ROOT / "charts"
REPORT_PATH = OUTPUT_ROOT / "SPX_NDX_IV_Robustness_2018_2024_Report.docx"


def shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    properties.append(element)


def set_cell(cell, value: object, *, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(value))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(8.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, frame: pd.DataFrame, caption: str) -> None:
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    caption_run.bold = True
    caption_run.font.name = "Times New Roman"
    caption_run.font.size = Pt(9.5)
    table = document.add_table(
        rows=1, cols=len(frame.columns), style="Table Grid"
    )
    for index, column in enumerate(frame.columns):
        set_cell(table.rows[0].cells[index], column, bold=True)
        shade(table.rows[0].cells[index], "D9EAF7")
    for row in frame.itertuples(index=False):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell(cells[index], value)
    document.add_paragraph()


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(6)


def add_equation(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(10.5)
    run.italic = True


def add_figure(
    document: Document,
    filename: str,
    caption: str,
    width: float = 6.5,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(
        str(CHART_ROOT / filename), width=Inches(width)
    )
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)


def build_report() -> None:
    summary = pd.read_csv(CSV_ROOT / "iv_robustness_summary.csv")
    filter_diagnostics = pd.read_csv(CSV_ROOT / "filter_diagnostics.csv")
    solver_diagnostics = pd.read_csv(CSV_ROOT / "solver_diagnostics.csv")
    yearly = pd.read_csv(CSV_ROOT / "yearly_iv_error_summary.csv")

    overall = summary[summary["ScopeType"] == "Overall"].iloc[0]
    market_rows = summary[summary["ScopeType"] == "Market"].set_index("Scope")
    tenor_rows = summary[summary["ScopeType"] == "Target tenor"].set_index("Scope")
    ndx_90 = summary[
        (summary["ScopeType"] == "Market × tenor")
        & (summary["Scope"] == "NDX — 90 days")
    ].iloc[0]

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = Pt(18)
    styles["Title"].font.bold = True
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Times New Roman"
        styles[style_name].font.color.rgb = RGBColor(31, 78, 120)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(
        "Independent Implied-Volatility Validation:\n"
        "SPX and NDX Options, 2018–2024"
    )
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Robustness appendix for the eSSVI/RND paper"
    ).italic = True

    document.add_heading("Executive Summary", level=1)
    add_body(
        document,
        (
            "This robustness exercise independently inverts bid–ask midpoint "
            "prices for SPX and NDX options while preserving the project's "
            "existing Black–Scholes and Newton/Brent solver. The independent "
            "sample does not require a valid OptionMetrics IV or OptionMetrics "
            "delta before inversion. It contains "
            f"{int(overall['NAttempted']):,} attempts, all of which converge. "
            "The paper comparison sample subsequently applies the main-sample "
            "absolute-delta restriction and intersects converged self IV with "
            f"valid OptionMetrics IV, yielding {int(overall['NCommonValid']):,} "
            "observations."
        ),
    )
    add_body(
        document,
        (
            "Agreement is extremely strong: correlation is "
            f"{overall['Correlation']:.6f}, MAE is "
            f"{100*overall['MAE']:.4f} volatility points, and RMSE is "
            f"{100*overall['RMSE']:.4f} points. Self-calculated IV is lower by "
            f"{100*overall['MeanErrorBias']:.4f} points on average. The bias is "
            "small economically but systematic, particularly for NDX. These "
            "results support OptionMetrics IV as the primary paper measure while "
            "also showing that exact replication depends on vendor conventions."
        ),
    )

    document.add_heading("1. What Was Changed", level=1)
    add_body(
        document,
        (
            "The Black–Scholes pricing functions, analytical Vega, no-arbitrage "
            "bounds, Newton iteration, multiple starting values, step-halving, "
            "and Brent fallback in part1_analysis.py were not rewritten. The "
            "single-date wrapper was replaced by a panel-data layer that supplies "
            "observation-specific spot, maturity, zero rate, dividend yield, and "
            "bid–ask midpoint."
        ),
    )
    add_equation(document, "Pmid = (Bid + Ask)/2")
    add_equation(
        document,
        "BS(St, K, T, rt,T, qt, σself) = Pmid",
    )
    add_body(
        document,
        (
            "Spot is the same-date OptionMetrics index close. T is actual calendar "
            "days to expiration divided by 365. The risk-free rate is linearly "
            "interpolated from the same-date OptionMetrics zero curve at the "
            "option's maturity, and q is the corresponding daily OptionMetrics "
            "index dividend yield. No constant market input is used across the "
            "2018–2024 panel."
        ),
    )

    document.add_heading("2. Sample Construction and Non-Circular Design", level=1)
    add_body(
        document,
        (
            "Month-end dates, target tenors, and chosen expirations exactly match "
            "the primary eSSVI analysis. Independent inversion requires valid "
            "positive bid and ask, ask above bid, relative spread no greater than "
            "50%, valid market inputs, OTM status under the independently "
            "constructed forward, and survival of duplicate resolution. It does "
            "not require OptionMetrics IV or delta."
        ),
    )
    add_body(
        document,
        (
            "After inversion, the comparison sample imposes the paper's "
            "0.05–0.95 absolute-delta restriction and requires both a valid "
            "OptionMetrics IV and a converged self IV. This sequence prevents "
            "OptionMetrics IV validity from determining whether an independent "
            "IV is attempted."
        ),
    )
    filter_display = filter_diagnostics.copy()
    filter_display["Count"] = filter_display["Count"].map(
        lambda value: f"{int(value):,}"
    )
    filter_display["Share (%)"] = (
        100 * filter_display["ShareWithinMarket"]
    ).map(lambda value: f"{value:.2f}")
    add_table(
        document,
        filter_display[
            ["Market", "FilterStatus", "Count", "Share (%)"]
        ],
        "Table 1. Independent-sample filter audit",
    )

    document.add_heading("3. Solver Diagnostics", level=1)
    add_body(
        document,
        (
            f"All {int(overall['NAttempted']):,} independent attempts converge. "
            "Newton solves the overwhelming majority, while Brent provides a "
            "successful fallback for difficult low-Vega or poorly conditioned "
            "quotes. No solver failure is hidden or removed."
        ),
    )
    solver_display = solver_diagnostics.copy()
    solver_display["Count"] = solver_display["Count"].map(
        lambda value: f"{int(value):,}"
    )
    add_table(
        document,
        solver_display[
            [
                "Market",
                "SolverStatus",
                "SolverMethod",
                "Converged",
                "Count",
            ]
        ],
        "Table 2. Newton/Brent convergence diagnostics",
    )

    document.add_heading("4. IV Comparison Results", level=1)
    summary_display = summary[
        summary["ScopeType"].isin(["Overall", "Market", "Target tenor"])
    ].copy()
    summary_display["N"] = summary_display["NCommonValid"].map(
        lambda value: f"{int(value):,}"
    )
    summary_display["Conv. (%)"] = (
        100 * summary_display["ConvergenceRate"]
    ).map(lambda value: f"{value:.2f}")
    for source, target in [
        ("MeanErrorBias", "Bias (pts)"),
        ("MAE", "MAE (pts)"),
        ("RMSE", "RMSE (pts)"),
        ("MedianAbsoluteError", "Median AE (pts)"),
    ]:
        summary_display[target] = (100 * summary_display[source]).map(
            lambda value: f"{value:.4f}"
        )
    summary_display["Correlation"] = summary_display["Correlation"].map(
        lambda value: f"{value:.6f}"
    )
    add_table(
        document,
        summary_display[
            [
                "Scope",
                "N",
                "Conv. (%)",
                "Bias (pts)",
                "MAE (pts)",
                "RMSE (pts)",
                "Median AE (pts)",
                "Correlation",
            ]
        ],
        "Table 3. Independent-IV robustness summary",
    )
    add_figure(
        document,
        "optionmetrics_vs_self_iv.png",
        (
            "Figure 1. OptionMetrics IV and self-calculated IV with the "
            "45-degree reference line."
        ),
    )

    document.add_heading("5. Discrepancies Across the Surface", level=1)
    add_body(
        document,
        (
            "The discrepancy is not uniform. SPX agreement is nearly exact at "
            "30 and 90 days, whereas NDX displays a modest negative bias at "
            "longer horizons and especially near the 90-day target. For NDX "
            f"90-day observations, MAE is {100*ndx_90['MAE']:.4f} volatility "
            f"points and bias is {100*ndx_90['MeanErrorBias']:.4f} points. "
            "Differences also become somewhat larger away from ATM. This pattern "
            "is reported rather than adjusted away."
        ),
    )
    add_body(
        document,
        (
            "Plausible explanations include OptionMetrics' proprietary forward "
            "and dividend treatment, maturity interpolation, exact settlement "
            "time, and quote synchronization. Because the errors are tiny relative "
            "to the volatility level and correlation remains above 0.9999, they "
            "do not overturn the paper's surface comparisons. They do caution "
            "against claiming an exact mechanical replication of vendor IV."
        ),
    )
    add_figure(
        document,
        "iv_difference_vs_log_moneyness.png",
        (
            "Figure 2. Self-calculated IV minus OptionMetrics IV across "
            "log-forward moneyness."
        ),
    )
    add_figure(
        document,
        "absolute_iv_error_by_tenor.png",
        "Figure 3. Absolute IV difference by market and target tenor.",
    )

    document.add_heading("6. Implication for the Main Paper", level=1)
    add_body(
        document,
        (
            "OptionMetrics IV should remain the primary input to eSSVI. The "
            "independent inversion strongly validates the vendor series but does "
            "not improve on its consistent treatment of settlement and forward "
            "conventions. The robustness output is maintained as a separate "
            "branch and does not overwrite the existing 651 eSSVI/RND slices, "
            "Excel workbook, figures, regressions, or empirical report."
        ),
    )

    document.add_heading("7. Assumptions and Limitations to Report", level=1)
    limitations = [
        "The option and index closes may not be perfectly synchronized.",
        "Zero rates are linearly interpolated and treated as continuously compounded.",
        "Dividend yields are daily OptionMetrics index yields and are held constant over each option's life.",
        "Time to maturity uses calendar days/365 and does not model intraday AM/PM settlement time.",
        "OptionMetrics' exact proprietary IV conventions are not fully observable.",
        "The common comparison sample applies the main paper's delta filter after independent inversion.",
        "Wing differences can be amplified by low Vega even when price discrepancies are small.",
    ]
    for item in limitations:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Output Traceability", level=1)
    add_body(
        document,
        (
            "The compressed row-level robustness dataset contains every eligible "
            "inversion and its market inputs, solver method, iterations, status, "
            "repricing error, and OptionMetrics comparison. A separate compressed "
            "row-level filter audit preserves removal reasons before inversion. "
            "All tables and figures in this report are generated from those files."
        ),
    )

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("SPX–NDX independent IV robustness report")
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)

    document.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    build_report()
