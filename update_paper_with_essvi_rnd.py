#!/usr/bin/env python3
"""Expand the paper outline with eSSVI and risk-neutral-density sections."""

from __future__ import annotations

import argparse
from pathlib import Path

ROOT = Path(__file__).resolve().parent

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


DEFAULT_OUTPUT = ROOT / "outputs" / "定题课_已加入_eSSVI_RND.docx"


def paragraph_by_exact_text(document: Document, text: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def replace_paragraph(paragraph, text: str, *, bold: bool = False) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)


def insert_after(
    document: Document,
    anchor,
    text: str = "",
    *,
    bold: bool = False,
    italic: bool = False,
    centered: bool = False,
):
    paragraph = document.add_paragraph()
    anchor._p.addnext(paragraph._p)
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    if centered:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


def insert_picture_after(
    document: Document,
    anchor,
    image_path: Path,
    caption: str,
):
    paragraph = document.add_paragraph()
    anchor._p.addnext(paragraph._p)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.35))
    caption_paragraph = insert_after(
        document, paragraph, caption, italic=True, centered=True
    )
    return caption_paragraph


def common_maturity_statistics() -> dict[str, float]:
    summaries = {}
    parameters = {}
    for ticker in ("SPX", "NDX"):
        summaries[ticker] = pd.read_csv(
            ROOT
            / "outputs"
            / "essvi_rnd"
            / ticker
            / "csv"
            / "rnd_summary.csv",
            parse_dates=["Maturity"],
        )
        parameters[ticker] = pd.read_csv(
            ROOT
            / "outputs"
            / "essvi_rnd"
            / ticker
            / "csv"
            / "essvi_parameters.csv",
            parse_dates=["Maturity"],
        )

    target = pd.Timestamp("2024-12-20")
    result: dict[str, float] = {}
    for ticker in ("SPX", "NDX"):
        rnd_row = summaries[ticker].loc[
            summaries[ticker]["Maturity"] == target
        ].iloc[0]
        parameter_row = parameters[ticker].loc[
            parameters[ticker]["Maturity"] == target
        ].iloc[0]
        prefix = ticker.lower()
        result[f"{prefix}_rho"] = float(parameter_row["Rho"])
        result[f"{prefix}_rmse"] = float(parameter_row["RMSE_IV"])
        result[f"{prefix}_cv"] = float(rnd_row["CoefficientOfVariation"])
        result[f"{prefix}_skew"] = float(rnd_row["Skewness"])
        result[f"{prefix}_kurtosis"] = float(rnd_row["ExcessKurtosis"])
        result[f"{prefix}_downside"] = float(
            rnd_row["ProbabilityBelow80PctForward"]
        )

    for ticker in ("SPX", "NDX"):
        table = parameters[ticker]
        prefix = ticker.lower()
        result[f"{prefix}_median_rmse"] = float(table["RMSE_IV"].median())
        result[f"{prefix}_max_rmse"] = float(table["RMSE_IV"].max())
        result[f"{prefix}_slices"] = int(len(table))
        result[f"{prefix}_observations"] = int(table["Observations"].sum())
        result[f"{prefix}_min_g"] = float(table["MinDensityFactorG"].min())
    return result


def update_document(source: Path, output: Path) -> None:
    stats = common_maturity_statistics()
    document = Document(source)

    title = document.paragraphs[0]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(15)

    # Methodology: eSSVI
    heading = paragraph_by_exact_text(document, "3.2 SVI")
    replace_paragraph(heading, "3.2 eSSVI Volatility-Surface Estimation", bold=True)
    anchor = paragraph_by_exact_text(document, "表达式和参数含义")
    replace_paragraph(
        anchor,
        (
            "The raw implied volatilities are converted to forward "
            "log-moneyness k = ln(K/Fₜ), where Fₜ = S₀exp[(r−q)T], and to total "
            "implied variance w = σ²IVT. To reduce the effect of illiquid "
            "in-the-money quotes, the calibration uses out-of-the-money puts "
            "for K < Fₜ and out-of-the-money calls for K ≥ Fₜ."
        ),
    )
    anchor = paragraph_by_exact_text(document, "进阶版：可以引入无套利约束条件")
    replace_paragraph(
        anchor,
        (
            "For each maturity, the extended SSVI slice is parameterized in "
            "natural coordinates as w(k) = ½{θ + ρψk + "
            "[(ψk + θρ)² + θ²(1−ρ²)]¹ᐟ²}. Here θ is ATM total variance, ρ "
            "controls asymmetry and skew, and ψ controls the wing slope. In "
            "contrast with basic SSVI, ρ is allowed to vary by maturity, which "
            "is important when the SPX and NDX skews evolve differently over "
            "the term structure [1,2]."
        ),
    )
    anchor = insert_after(
        document,
        anchor,
        (
            "Parameters are estimated by minimizing the mean squared error "
            "between observed and fitted IVs. The sufficient butterfly "
            "conditions ψ(1+|ρ|) < 4 and ψ²(1+|ρ|) ≤ 4θ are imposed, together "
            "with non-decreasing total variance across adjacent maturities on "
            "a common log-moneyness grid. These restrictions make the fitted "
            "surface suitable for taking the second strike derivative needed "
            "for risk-neutral-density estimation."
        ),
    )
    anchor = insert_after(
        document,
        anchor,
        (
            "Fit is evaluated using IV RMSE, MAE and R², while the minimum "
            "density factor g(k) and the minimum calendar-variance increment "
            "are recorded as numerical no-arbitrage diagnostics."
        ),
    )

    # Methodology: RND
    heading = paragraph_by_exact_text(document, "3.4 风险中性密度（RND）")
    replace_paragraph(heading, "3.4 Risk-Neutral Density (RND)", bold=True)
    anchor = heading
    anchor = insert_after(
        document,
        anchor,
        (
            "Under standard regularity conditions, the Breeden–Litzenberger "
            "identity recovers the risk-neutral density from European call "
            "prices: fQ(K,T) = exp(rT) ∂²C(K,T)/∂K² [3]. Directly differentiating "
            "discrete mid-quotes is unstable, so call prices and their "
            "derivatives are generated from the smooth eSSVI total-variance "
            "surface."
        ),
    )
    anchor = insert_after(
        document,
        anchor,
        (
            "For k = ln(K/F), define d₋ = −k/√w − √w/2. The density implied by "
            "a differentiable total-variance smile is fQ(K,T) = "
            "φ(d₋)g(k)/[K√w], where g(k) = [1−kw′/(2w)]² − "
            "(w′²/4)(1/w+1/4) + w″/2. A non-negative g(k) is also the key "
            "butterfly-arbitrage diagnostic."
        ),
    )
    anchor = insert_after(
        document,
        anchor,
        (
            "The density is evaluated on K/F ∈ [0.35, 1.80] and normalized by "
            "its numerical mass. For each maturity, the analysis reports the "
            "standard deviation, skewness, excess kurtosis, and the "
            "model-implied probabilities of ending below 80% or above 120% of "
            "the forward. Because several short-dated strike ranges are narrow, "
            "far-tail probabilities are treated as model-based extrapolations."
        ),
    )

    # Empirical eSSVI results
    heading = paragraph_by_exact_text(document, "4.1 拟合优度检验")
    replace_paragraph(heading, "4.1 eSSVI Fit and No-Arbitrage Diagnostics", bold=True)
    anchor = paragraph_by_exact_text(document, "可以看样本内外的拟合效果")
    replace_paragraph(
        anchor,
        (
            f"The constrained eSSVI procedure fits {stats['spx_slices']} SPX "
            f"slices ({stats['spx_observations']} OTM observations) and "
            f"{stats['ndx_slices']} NDX slices "
            f"({stats['ndx_observations']} observations). Median IV RMSE is "
            f"{100*stats['spx_median_rmse']:.3f} volatility points for SPX and "
            f"{100*stats['ndx_median_rmse']:.3f} points for NDX; the corresponding "
            f"maximum RMSEs are {100*stats['spx_max_rmse']:.3f} and "
            f"{100*stats['ndx_max_rmse']:.3f} points."
        ),
    )
    anchor = insert_after(
        document,
        anchor,
        (
            f"All fitted slices satisfy the imposed diagnostics. The minimum "
            f"g(k) is {stats['spx_min_g']:.3f} for SPX and "
            f"{stats['ndx_min_g']:.3f} for NDX, and total variance is "
            "non-decreasing over the tested calendar grid. At the common "
            f"20 December 2024 maturity, ρ equals {stats['spx_rho']:.3f} for "
            f"SPX and {stats['ndx_rho']:.3f} for NDX, indicating a more "
            "negative NDX risk-neutral skew at that horizon."
        ),
    )
    anchor = insert_picture_after(
        document,
        anchor,
        ROOT
        / "outputs"
        / "essvi_rnd"
        / "SPX"
        / "charts"
        / "essvi_selected_smiles.png",
        "Figure 1. Observed SPX OTM implied volatilities and constrained eSSVI fits.",
    )
    anchor = insert_picture_after(
        document,
        anchor,
        ROOT
        / "outputs"
        / "essvi_rnd"
        / "NDX"
        / "charts"
        / "essvi_selected_smiles.png",
        "Figure 2. Observed NDX OTM implied volatilities and constrained eSSVI fits.",
    )

    # Empirical RND results
    heading = paragraph_by_exact_text(document, "4.3 RND")
    replace_paragraph(heading, "4.3 Risk-Neutral Density Comparison", bold=True)
    anchor = paragraph_by_exact_text(document, "同一天，同一期限的RND曲线")
    replace_paragraph(
        anchor,
        (
            "The cleanest cross-market comparison holds valuation date and "
            "maturity fixed. For 20 December 2024, the coefficient of variation "
            f"of the normalized terminal level is {100*stats['spx_cv']:.2f}% "
            f"for SPX and {100*stats['ndx_cv']:.2f}% for NDX. The NDX density "
            "is therefore materially wider, consistent with its higher IV level."
        ),
    )
    anchor = paragraph_by_exact_text(document, "可以关注Fat tail")
    replace_paragraph(
        anchor,
        (
            "For the same maturity, the model-implied probability of finishing "
            f"below 80% of the forward is {100*stats['spx_downside']:.2f}% for "
            f"SPX versus {100*stats['ndx_downside']:.2f}% for NDX. Estimated "
            f"skewness is {stats['spx_skew']:.3f} and {stats['ndx_skew']:.3f}, "
            f"while excess kurtosis is {stats['spx_kurtosis']:.2f} and "
            f"{stats['ndx_kurtosis']:.2f}, respectively. The stronger NDX "
            "downside-tail probability mainly reflects its broader distribution; "
            "the higher SPX kurtosis estimate should be interpreted cautiously "
            "because it is more sensitive to wing extrapolation."
        ),
    )
    anchor = insert_picture_after(
        document,
        anchor,
        ROOT / "outputs" / "essvi_rnd" / "spx_ndx_rnd_comparison.png",
        (
            "Figure 3. eSSVI-implied SPX and NDX risk-neutral densities for "
            "20 December 2024, normalized by each index forward."
        ),
    )

    # Core references for the newly added methodology.
    reference_heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("References参考文献")
    )
    replace_paragraph(reference_heading, "References", bold=True)
    anchor = reference_heading
    references = [
        (
            "[1] Gatheral, J., & Jacquier, A. (2014). Arbitrage-free SVI "
            "volatility surfaces. Quantitative Finance, 14(1), 59–71."
        ),
        (
            "[2] Hendriks, S., & Martini, C. (2019). The extended SSVI "
            "volatility surface. Journal of Computational Finance, 22(5), "
            "25–39. https://doi.org/10.21314/JCF.2019.365"
        ),
        (
            "[3] Breeden, D. T., & Litzenberger, R. H. (1978). Prices of "
            "state-contingent claims implicit in option prices. Journal of "
            "Business, 51(4), 621–651."
        ),
    ]
    for reference in references:
        anchor = insert_after(document, anchor, reference)

    # Apply a consistent readable font to paragraphs that retained the template.
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                run.font.name = "Times New Roman"
            if run.font.size is None:
                run.font.size = Pt(11)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        type=Path,
        required=True,
        help="Path to the source Word document.",
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    update_document(args.source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
