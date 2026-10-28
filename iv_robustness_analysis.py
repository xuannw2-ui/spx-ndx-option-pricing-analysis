#!/usr/bin/env python3
"""Independent IV validation for the 2018--2024 SPX/NDX panel.

This module intentionally reuses the Black--Scholes and Newton/Brent inversion
implemented in ``part1_analysis.py``.  It does not feed self-calculated IV back
into the primary eSSVI pipeline; OptionMetrics IV remains the main measure.
"""

from __future__ import annotations

import argparse
import math
import warnings
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from long_horizon_essvi_analysis import (
    SECID_TO_MARKET,
    YEARS,
    configure_data_root,
    interpolate_zero_rate,
    load_auxiliary_data,
    option_path,
    security_path,
)
from part1_analysis import implied_volatility_newton


DEFAULT_OUTPUT = Path("outputs/iv_robustness_2018_2024")
MAIN_METRICS = Path(
    "outputs/long_horizon_essvi/csv/monthly_essvi_rnd_metrics.csv"
)

RAW_COLUMNS = [
    "secid",
    "date",
    "exdate",
    "cp_flag",
    "strike_price",
    "best_bid",
    "best_offer",
    "volume",
    "open_interest",
    "impl_volatility",
    "delta",
    "symbol",
    "ss_flag",
    "am_settlement",
]

ROBUSTNESS_COLUMNS = [
    "Date",
    "Market",
    "Expiration",
    "DaysToMaturity",
    "TargetTenorDays",
    "OptionType",
    "Strike",
    "Spot",
    "Forward",
    "LogForwardMoneyness",
    "Bid",
    "Ask",
    "Midpoint",
    "RiskFreeRate",
    "DividendYield",
    "OptionMetricsIV",
    "OptionMetricsIVValid",
    "MainSampleDeltaEligible",
    "SelfCalculatedIV",
    "IVDifference",
    "AbsoluteIVDifference",
    "BSRepricedValue",
    "PriceError",
    "SolverMethod",
    "SolverIterations",
    "Converged",
    "SolverStatus",
    "FilterStatus",
    "RelativeBidAskSpread",
    "Delta",
    "Volume",
    "OpenInterest",
    "SecID",
    "Symbol",
    "SpecialSettlementFlag",
    "AMSettlementFlag",
]


def load_main_slice_choices(path: Path = MAIN_METRICS) -> pd.DataFrame:
    """Use exactly the date/expiration/target-tenor choices of the main paper."""
    metrics = pd.read_csv(
        path,
        usecols=[
            "Date",
            "Market",
            "Maturity",
            "TargetTenorDays",
            "ActualDaysToMaturity",
        ],
        parse_dates=["Date", "Maturity"],
    )
    market_to_secid = {market: secid for secid, market in SECID_TO_MARKET.items()}
    metrics["secid"] = metrics["Market"].map(market_to_secid)
    return (
        metrics.rename(
            columns={"Date": "date", "Maturity": "exdate"}
        )[
            [
                "date",
                "secid",
                "exdate",
                "TargetTenorDays",
                "ActualDaysToMaturity",
            ]
        ]
        .drop_duplicates()
        .sort_values(["date", "secid", "TargetTenorDays"])
        .reset_index(drop=True)
    )


def assign_first_failure(
    frame: pd.DataFrame,
    condition: pd.Series,
    status: str,
) -> None:
    mask = frame["FilterStatus"].eq("eligible") & condition.fillna(False)
    frame.loc[mask, "FilterStatus"] = status


def construct_market_inputs(
    observations: pd.DataFrame,
    security: pd.DataFrame,
    dividend_lookup: dict[tuple[int, pd.Timestamp], float],
    zero_by_date: dict[pd.Timestamp, tuple[np.ndarray, np.ndarray]],
) -> pd.DataFrame:
    """Attach observation-specific S, r(T), q, T, and forward values."""
    spot = security.rename(columns={"close": "Spot"})[
        ["secid", "date", "Spot"]
    ]
    result = observations.merge(spot, on=["secid", "date"], how="left")
    result["DaysToMaturity"] = (result["exdate"] - result["date"]).dt.days
    result["TimeToMaturityYears"] = result["DaysToMaturity"] / 365.0
    result["RiskFreeRate"] = [
        interpolate_zero_rate(zero_by_date, date, days)
        for date, days in zip(result["date"], result["DaysToMaturity"])
    ]
    result["DividendYield"] = [
        dividend_lookup.get((int(secid), pd.Timestamp(date)), math.nan)
        for secid, date in zip(result["secid"], result["date"])
    ]
    result["Forward"] = result["Spot"] * np.exp(
        (result["RiskFreeRate"] - result["DividendYield"])
        * result["TimeToMaturityYears"]
    )
    result["Strike"] = result["strike_price"] / 1000.0
    result["LogForwardMoneyness"] = np.log(
        result["Strike"] / result["Forward"]
    )
    return result


def classify_and_deduplicate(observations: pd.DataFrame) -> pd.DataFrame:
    """Apply quote filters without requiring a valid OptionMetrics IV."""
    frame = observations.copy()
    frame["Midpoint"] = 0.5 * (frame["best_bid"] + frame["best_offer"])
    frame["RelativeBidAskSpread"] = (
        (frame["best_offer"] - frame["best_bid"]) / frame["Midpoint"]
    )
    frame["MainSampleDeltaEligible"] = (
        frame["delta"].notna() & frame["delta"].abs().between(0.05, 0.95)
    )
    frame["FilterStatus"] = "eligible"

    assign_first_failure(
        frame,
        frame["best_bid"].isna() | frame["best_offer"].isna(),
        "missing_bid_or_ask",
    )
    assign_first_failure(
        frame,
        (frame["best_bid"] <= 0)
        | (frame["best_offer"] <= 0)
        | (frame["best_offer"] <= frame["best_bid"])
        | (frame["Midpoint"] <= 0),
        "invalid_bid_ask",
    )
    assign_first_failure(
        frame,
        ~frame["RelativeBidAskSpread"].between(0, 0.50),
        "relative_spread_above_50pct",
    )
    assign_first_failure(
        frame,
        frame["Strike"].isna() | (frame["Strike"] <= 0),
        "invalid_strike",
    )
    assign_first_failure(
        frame,
        frame[
            [
                "Spot",
                "RiskFreeRate",
                "DividendYield",
                "TimeToMaturityYears",
                "Forward",
            ]
        ]
        .isna()
        .any(axis=1)
        | (frame["Spot"] <= 0)
        | (frame["TimeToMaturityYears"] <= 0)
        | (frame["Forward"] <= 0),
        "missing_or_invalid_market_input",
    )
    is_otm = (
        (frame["cp_flag"] == "P") & (frame["Strike"] < frame["Forward"])
    ) | (
        (frame["cp_flag"] == "C") & (frame["Strike"] >= frame["Forward"])
    )
    assign_first_failure(frame, ~is_otm, "outside_otm_smile")

    eligible = frame[frame["FilterStatus"] == "eligible"].sort_values(
        [
            "date",
            "secid",
            "TargetTenorDays",
            "cp_flag",
            "Strike",
            "RelativeBidAskSpread",
            "open_interest",
        ],
        ascending=[True, True, True, True, True, True, False],
    )
    duplicate = eligible.duplicated(
        [
            "date",
            "secid",
            "TargetTenorDays",
            "cp_flag",
            "Strike",
        ],
        keep="first",
    )
    frame.loc[eligible.index[duplicate], "FilterStatus"] = (
        "duplicate_less_preferred"
    )
    frame["OptionMetricsIVValid"] = frame["impl_volatility"].between(
        0.01, 3.00
    ) & frame["impl_volatility"].notna()
    return frame


def extract_independent_sample(
    output_root: Path,
    *,
    force: bool = False,
) -> pd.DataFrame:
    """Stream raw annual files and preserve every row at selected expirations."""
    cache = output_root / "cache" / "classified_observations.pkl.gz"
    if cache.exists() and not force:
        return pd.read_pickle(cache, compression="gzip")

    choices = load_main_slice_choices()
    dividends, zero_curve, _ = load_auxiliary_data()
    dividend_lookup = (
        dividends.set_index(["secid", "date"])["DividendYield"].to_dict()
    )
    zero_by_date = {
        pd.Timestamp(date): (
            group["days"].to_numpy(dtype=float),
            group["ZeroRate"].to_numpy(dtype=float),
        )
        for date, group in zero_curve.sort_values(["date", "days"]).groupby("date")
    }

    yearly_frames: list[pd.DataFrame] = []
    for year in YEARS:
        year_choices = choices[choices["date"].dt.year == year].copy()
        year_choice_strings = set(
            year_choices["date"].dt.strftime("%Y-%m-%d")
        )
        merge_choices = year_choices[
            ["date", "secid", "exdate", "TargetTenorDays"]
        ]
        retained: list[pd.DataFrame] = []
        for chunk in pd.read_csv(
            option_path(year),
            usecols=RAW_COLUMNS,
            chunksize=500_000,
            low_memory=False,
        ):
            chunk = chunk[chunk["date"].isin(year_choice_strings)]
            if chunk.empty:
                continue
            chunk["date"] = pd.to_datetime(chunk["date"])
            chunk["exdate"] = pd.to_datetime(chunk["exdate"])
            chunk = chunk.merge(
                merge_choices,
                on=["date", "secid", "exdate"],
                how="inner",
            )
            if not chunk.empty:
                retained.append(chunk)
        if not retained:
            continue
        observations = pd.concat(retained, ignore_index=True)
        numeric = [
            "strike_price",
            "best_bid",
            "best_offer",
            "volume",
            "open_interest",
            "impl_volatility",
            "delta",
            "ss_flag",
            "am_settlement",
        ]
        for column in numeric:
            observations[column] = pd.to_numeric(
                observations[column], errors="coerce"
            )
        security = pd.read_csv(security_path(year), parse_dates=["date"])
        observations = construct_market_inputs(
            observations,
            security,
            dividend_lookup,
            zero_by_date,
        )
        observations["Market"] = observations["secid"].map(SECID_TO_MARKET)
        observations = classify_and_deduplicate(observations)
        yearly_frames.append(observations)
        print(
            f"extracted {year}: {len(observations):,} selected-expiration rows, "
            f"{(observations['FilterStatus'] == 'eligible').sum():,} eligible",
            flush=True,
        )

    classified = pd.concat(yearly_frames, ignore_index=True)
    cache.parent.mkdir(parents=True, exist_ok=True)
    classified.to_pickle(cache, compression="gzip")
    return classified


def calculate_independent_iv(
    classified: pd.DataFrame,
    output_root: Path,
    *,
    force: bool = False,
) -> pd.DataFrame:
    """Run the unchanged Newton/Brent solver on every independently eligible row."""
    cache = output_root / "cache" / "iv_robustness_observations.pkl.gz"
    if cache.exists() and not force:
        return pd.read_pickle(cache, compression="gzip")

    eligible = classified[classified["FilterStatus"] == "eligible"].copy()
    self_iv = np.full(len(eligible), np.nan)
    repriced = np.full(len(eligible), np.nan)
    price_error = np.full(len(eligible), np.nan)
    iterations = np.zeros(len(eligible), dtype=int)
    methods = np.full(len(eligible), "", dtype=object)
    converged = np.zeros(len(eligible), dtype=bool)
    statuses = np.full(len(eligible), "", dtype=object)

    for position, row in enumerate(eligible.itertuples(index=False)):
        result = implied_volatility_newton(
            option_type="Call" if row.cp_flag == "C" else "Put",
            market_price=float(row.Midpoint),
            spot=float(row.Spot),
            strike=float(row.Strike),
            rate=float(row.RiskFreeRate),
            dividend_yield=float(row.DividendYield),
            time_to_maturity=float(row.TimeToMaturityYears),
        )
        self_iv[position] = float(result["iv"])
        repriced[position] = float(result["repriced"])
        price_error[position] = float(result["error"])
        iterations[position] = int(result["iterations"])
        methods[position] = str(result["method"])
        converged[position] = bool(result["converged"])
        statuses[position] = str(result["status"])
        if (position + 1) % 20_000 == 0:
            print(
                f"solved {position + 1:,}/{len(eligible):,} eligible quotes",
                flush=True,
            )

    eligible["SelfCalculatedIV"] = self_iv
    eligible["BSRepricedValue"] = repriced
    eligible["PriceError"] = price_error
    eligible["SolverIterations"] = iterations
    eligible["SolverMethod"] = methods
    eligible["Converged"] = converged
    eligible["SolverStatus"] = statuses
    eligible["IVDifference"] = (
        eligible["SelfCalculatedIV"] - eligible["impl_volatility"]
    )
    eligible["AbsoluteIVDifference"] = eligible["IVDifference"].abs()
    eligible.to_pickle(cache, compression="gzip")
    return eligible


def format_robustness_dataset(eligible: pd.DataFrame) -> pd.DataFrame:
    output = pd.DataFrame(
        {
            "Date": eligible["date"],
            "Market": eligible["Market"],
            "Expiration": eligible["exdate"],
            "DaysToMaturity": eligible["DaysToMaturity"],
            "TargetTenorDays": eligible["TargetTenorDays"],
            "OptionType": eligible["cp_flag"].map({"C": "Call", "P": "Put"}),
            "Strike": eligible["Strike"],
            "Spot": eligible["Spot"],
            "Forward": eligible["Forward"],
            "LogForwardMoneyness": eligible["LogForwardMoneyness"],
            "Bid": eligible["best_bid"],
            "Ask": eligible["best_offer"],
            "Midpoint": eligible["Midpoint"],
            "RiskFreeRate": eligible["RiskFreeRate"],
            "DividendYield": eligible["DividendYield"],
            "OptionMetricsIV": eligible["impl_volatility"],
            "OptionMetricsIVValid": eligible["OptionMetricsIVValid"],
            "MainSampleDeltaEligible": eligible["MainSampleDeltaEligible"],
            "SelfCalculatedIV": eligible["SelfCalculatedIV"],
            "IVDifference": eligible["IVDifference"],
            "AbsoluteIVDifference": eligible["AbsoluteIVDifference"],
            "BSRepricedValue": eligible["BSRepricedValue"],
            "PriceError": eligible["PriceError"],
            "SolverMethod": eligible["SolverMethod"],
            "SolverIterations": eligible["SolverIterations"],
            "Converged": eligible["Converged"],
            "SolverStatus": eligible["SolverStatus"],
            "FilterStatus": eligible["FilterStatus"],
            "RelativeBidAskSpread": eligible["RelativeBidAskSpread"],
            "Delta": eligible["delta"],
            "Volume": eligible["volume"],
            "OpenInterest": eligible["open_interest"],
            "SecID": eligible["secid"],
            "Symbol": eligible["symbol"],
            "SpecialSettlementFlag": eligible["ss_flag"],
            "AMSettlementFlag": eligible["am_settlement"],
        }
    )
    return output[ROBUSTNESS_COLUMNS]


def summary_statistics(
    dataset: pd.DataFrame,
    *,
    scope_type: str,
    scope: str,
) -> dict[str, object]:
    self_valid = dataset["Converged"] & dataset["SelfCalculatedIV"].notna()
    common = dataset[
        self_valid
        & dataset["OptionMetricsIVValid"]
        & dataset["MainSampleDeltaEligible"]
    ]
    differences = common["IVDifference"]
    correlation = (
        common[["OptionMetricsIV", "SelfCalculatedIV"]]
        .corr()
        .iloc[0, 1]
        if len(common) >= 2
        else math.nan
    )
    return {
        "ScopeType": scope_type,
        "Scope": scope,
        "NAttempted": len(dataset),
        "NSelfIVConverged": int(self_valid.sum()),
        "ConvergenceRate": float(self_valid.mean()) if len(dataset) else math.nan,
        "NOptionMetricsIVValid": int(dataset["OptionMetricsIVValid"].sum()),
        "NCommonValid": len(common),
        "MeanErrorBias": float(differences.mean()),
        "MAE": float(differences.abs().mean()),
        "RMSE": float(np.sqrt(np.mean(differences**2))),
        "MedianAbsoluteError": float(differences.abs().median()),
        "Correlation": float(correlation),
    }


def build_summary_tables(
    dataset: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    rows = [summary_statistics(dataset, scope_type="Overall", scope="Full sample")]
    for market, group in dataset.groupby("Market"):
        rows.append(summary_statistics(group, scope_type="Market", scope=market))
    for tenor, group in dataset.groupby("TargetTenorDays"):
        rows.append(
            summary_statistics(
                group, scope_type="Target tenor", scope=f"{int(tenor)} days"
            )
        )
    for (market, tenor), group in dataset.groupby(
        ["Market", "TargetTenorDays"]
    ):
        rows.append(
            summary_statistics(
                group,
                scope_type="Market × tenor",
                scope=f"{market} — {int(tenor)} days",
            )
        )
    summary = pd.DataFrame(rows)

    common = dataset[
        dataset["Converged"]
        & dataset["OptionMetricsIVValid"]
        & dataset["MainSampleDeltaEligible"]
    ].copy()
    bucket_edges = [
        -np.inf,
        -0.20,
        -0.10,
        -0.05,
        0.0,
        0.05,
        0.10,
        0.20,
        np.inf,
    ]
    bucket_labels = [
        "<−0.20",
        "−0.20 to −0.10",
        "−0.10 to −0.05",
        "−0.05 to 0",
        "0 to 0.05",
        "0.05 to 0.10",
        "0.10 to 0.20",
        ">0.20",
    ]
    common["MoneynessBucket"] = pd.cut(
        common["LogForwardMoneyness"],
        bins=bucket_edges,
        labels=bucket_labels,
        include_lowest=True,
        right=False,
    )
    moneyness = (
        common.groupby(
            ["Market", "TargetTenorDays", "MoneynessBucket"],
            observed=True,
        )
        .agg(
            N=("IVDifference", "size"),
            MeanErrorBias=("IVDifference", "mean"),
            MAE=("AbsoluteIVDifference", "mean"),
            RMSE=(
                "IVDifference",
                lambda values: float(np.sqrt(np.mean(values**2))),
            ),
            MedianAbsoluteError=("AbsoluteIVDifference", "median"),
        )
        .reset_index()
    )
    yearly = (
        common.assign(Year=common["Date"].dt.year)
        .groupby(["Year", "Market"])
        .agg(
            N=("IVDifference", "size"),
            MeanErrorBias=("IVDifference", "mean"),
            MAE=("AbsoluteIVDifference", "mean"),
            RMSE=(
                "IVDifference",
                lambda values: float(np.sqrt(np.mean(values**2))),
            ),
            Correlation=(
                "SelfCalculatedIV",
                lambda values: math.nan,
            ),
        )
        .reset_index()
    )
    correlations = (
        common.assign(Year=common["Date"].dt.year)
        .groupby(["Year", "Market"])
        .apply(
            lambda group: group["OptionMetricsIV"].corr(
                group["SelfCalculatedIV"]
            ),
            include_groups=False,
        )
        .rename("Correlation")
        .reset_index()
    )
    yearly = yearly.drop(columns=["Correlation"]).merge(
        correlations, on=["Year", "Market"], how="left"
    )
    return summary, moneyness, yearly


def build_diagnostics(
    classified: pd.DataFrame,
    dataset: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    observed_filter_counts = (
        classified.groupby(["Market", "FilterStatus"], dropna=False)
        .size()
        .rename("Count")
    )
    expected_filter_statuses = [
        "eligible",
        "missing_bid_or_ask",
        "invalid_bid_ask",
        "relative_spread_above_50pct",
        "invalid_strike",
        "missing_or_invalid_market_input",
        "outside_otm_smile",
        "duplicate_less_preferred",
    ]
    complete_index = pd.MultiIndex.from_product(
        [["SPX", "NDX"], expected_filter_statuses],
        names=["Market", "FilterStatus"],
    )
    filter_counts = (
        observed_filter_counts.reindex(complete_index, fill_value=0)
        .reset_index()
    )
    filter_counts["ShareWithinMarket"] = filter_counts["Count"] / filter_counts.groupby(
        "Market"
    )["Count"].transform("sum")
    solver_counts = (
        dataset.groupby(
            ["Market", "SolverStatus", "SolverMethod", "Converged"],
            dropna=False,
        )
        .size()
        .rename("Count")
        .reset_index()
    )
    annual_filter = (
        classified.assign(Year=classified["date"].dt.year)
        .groupby(["Year", "Market", "FilterStatus"])
        .size()
        .rename("Count")
        .reset_index()
    )
    comparison_eligibility = (
        dataset.groupby(
            [
                "Market",
                "MainSampleDeltaEligible",
                "OptionMetricsIVValid",
                "Converged",
            ],
            dropna=False,
        )
        .size()
        .rename("Count")
        .reset_index()
    )
    return filter_counts, solver_counts, annual_filter, comparison_eligibility


def diagnostic_codebook() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ("FilterStatus", "eligible", "Passed independent pre-inversion filters."),
            ("FilterStatus", "missing_bid_or_ask", "Bid or ask is missing."),
            ("FilterStatus", "invalid_bid_ask", "Non-positive or incorrectly ordered bid/ask."),
            (
                "FilterStatus",
                "relative_spread_above_50pct",
                "Relative bid-ask spread exceeds 50%.",
            ),
            ("FilterStatus", "invalid_strike", "Strike is missing or non-positive."),
            (
                "FilterStatus",
                "missing_or_invalid_market_input",
                "Spot, maturity, zero rate, dividend yield, or forward is unavailable/invalid.",
            ),
            (
                "FilterStatus",
                "outside_otm_smile",
                "Quote is not an OTM put below forward or OTM call at/above forward.",
            ),
            (
                "FilterStatus",
                "duplicate_less_preferred",
                "Duplicate strike/type quote with a wider spread or lower open interest.",
            ),
            ("SolverStatus", "converged", "Newton or Brent inversion converged."),
            (
                "SolverStatus",
                "outside_no_arbitrage_bounds",
                "Midpoint violates the Black-Scholes no-arbitrage price bounds.",
            ),
            (
                "SolverStatus",
                "outside_volatility_bounds",
                "No IV exists within the permitted volatility interval.",
            ),
            (
                "SolverStatus",
                "brent_root_not_bracketed",
                "Brent fallback could not bracket an IV root.",
            ),
            ("SolverStatus", "brent_failed", "Brent fallback failed numerically."),
            (
                "SolverStatus",
                "newton_and_brent_not_converged",
                "Neither Newton nor Brent met the price tolerance.",
            ),
            ("SolverStatus", "invalid_input", "A required numerical input is invalid."),
            (
                "SolverStatus",
                "non_finite_input",
                "A required numerical input is non-finite.",
            ),
            (
                "SolverMethod",
                "Brent",
                "Legacy solver used its Brent fallback; this includes low-Vega or unsuccessful-Newton cases.",
            ),
        ],
        columns=["DiagnosticField", "Code", "Meaning"],
    )


def plot_iv_comparison(dataset: pd.DataFrame, chart_dir: Path) -> None:
    chart_dir.mkdir(parents=True, exist_ok=True)
    common = dataset[
        dataset["Converged"]
        & dataset["OptionMetricsIVValid"]
        & dataset["MainSampleDeltaEligible"]
    ].copy()
    lower = float(
        min(common["OptionMetricsIV"].quantile(0.001), common["SelfCalculatedIV"].quantile(0.001))
    )
    upper = float(
        max(common["OptionMetricsIV"].quantile(0.999), common["SelfCalculatedIV"].quantile(0.999))
    )

    fig, axes = plt.subplots(1, 2, figsize=(13, 6), sharex=True, sharey=True)
    for ax, market in zip(axes, ("SPX", "NDX")):
        group = common[common["Market"] == market]
        image = ax.hexbin(
            100 * group["OptionMetricsIV"],
            100 * group["SelfCalculatedIV"],
            gridsize=65,
            mincnt=1,
            bins="log",
            cmap="viridis",
        )
        ax.plot(
            [100 * lower, 100 * upper],
            [100 * lower, 100 * upper],
            color="red",
            ls="--",
            lw=1.5,
            label="45° line",
        )
        ax.set_title(f"{market} (N={len(group):,})")
        ax.set_xlabel("OptionMetrics IV (%)")
        ax.grid(alpha=0.2)
        ax.legend()
    axes[0].set_ylabel("Self-calculated IV (%)")
    fig.colorbar(image, ax=axes, label="log observation count")
    fig.suptitle("OptionMetrics IV versus independently calculated IV")
    fig.savefig(
        chart_dir / "optionmetrics_vs_self_iv.png",
        dpi=220,
        bbox_inches="tight",
    )
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(14, 6), sharex=True, sharey=True)
    rng = np.random.default_rng(401)
    for ax, market in zip(axes, ("SPX", "NDX")):
        group = common[common["Market"] == market].copy()
        plotted = (
            group.sample(n=min(25_000, len(group)), random_state=401)
            if len(group) > 25_000
            else group
        )
        ax.scatter(
            plotted["LogForwardMoneyness"],
            100 * plotted["IVDifference"],
            s=4,
            alpha=0.12,
            color="#4C78A8",
            rasterized=True,
        )
        bins = np.linspace(
            group["LogForwardMoneyness"].quantile(0.005),
            group["LogForwardMoneyness"].quantile(0.995),
            31,
        )
        group["PlotBucket"] = pd.cut(
            group["LogForwardMoneyness"], bins=bins, include_lowest=True
        )
        binned = (
            group.groupby("PlotBucket", observed=True)
            .agg(
                X=("LogForwardMoneyness", "mean"),
                Mean=("IVDifference", "mean"),
                Lower=("IVDifference", lambda values: values.quantile(0.10)),
                Upper=("IVDifference", lambda values: values.quantile(0.90)),
            )
            .dropna()
        )
        ax.plot(
            binned["X"],
            100 * binned["Mean"],
            color="#D62728",
            lw=2,
            label="Bin mean",
        )
        ax.fill_between(
            binned["X"].to_numpy(dtype=float),
            100 * binned["Lower"].to_numpy(dtype=float),
            100 * binned["Upper"].to_numpy(dtype=float),
            color="#D62728",
            alpha=0.15,
            label="10th–90th percentile",
        )
        ax.axhline(0, color="black", ls="--", lw=1)
        ax.set_title(market)
        ax.set_xlabel("Log-forward moneyness")
        ax.grid(alpha=0.2)
        ax.legend()
    axes[0].set_ylabel("Self IV − OptionMetrics IV (volatility points)")
    fig.suptitle("IV discrepancy across the smile")
    fig.tight_layout()
    fig.savefig(
        chart_dir / "iv_difference_vs_log_moneyness.png",
        dpi=220,
        bbox_inches="tight",
    )
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(13, 5.5), sharey=True)
    for ax, market in zip(axes, ("SPX", "NDX")):
        group = common[common["Market"] == market]
        data = [
            100
            * group.loc[
                group["TargetTenorDays"] == tenor,
                "AbsoluteIVDifference",
            ].to_numpy()
            for tenor in (30, 90, 180, 360)
        ]
        ax.boxplot(
            data,
            tick_labels=["30", "90", "180", "360"],
            showfliers=False,
        )
        ax.set_title(market)
        ax.set_xlabel("Target tenor (days)")
        ax.grid(axis="y", alpha=0.2)
    axes[0].set_ylabel("Absolute IV difference (volatility points)")
    fig.suptitle("Absolute IV discrepancy by market and target tenor")
    fig.tight_layout()
    fig.savefig(
        chart_dir / "absolute_iv_error_by_tenor.png",
        dpi=220,
        bbox_inches="tight",
    )
    plt.close(fig)


def write_workbook(
    summary: pd.DataFrame,
    moneyness: pd.DataFrame,
    yearly: pd.DataFrame,
    filter_counts: pd.DataFrame,
    solver_counts: pd.DataFrame,
    annual_filter: pd.DataFrame,
    comparison_eligibility: pd.DataFrame,
    codebook: pd.DataFrame,
    dataset: pd.DataFrame,
    output_path: Path,
) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="IV Summary", index=False)
        moneyness.to_excel(writer, sheet_name="Moneyness Summary", index=False)
        yearly.to_excel(writer, sheet_name="Yearly Summary", index=False)
        filter_counts.to_excel(writer, sheet_name="Filter Diagnostics", index=False)
        solver_counts.to_excel(writer, sheet_name="Solver Diagnostics", index=False)
        annual_filter.to_excel(writer, sheet_name="Annual Filters", index=False)
        comparison_eligibility.to_excel(
            writer, sheet_name="Comparison Eligibility", index=False
        )
        codebook.to_excel(writer, sheet_name="Diagnostic Codebook", index=False)
        dataset.to_excel(writer, sheet_name="Robustness Observations", index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(output_path)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center")
        for index, cells in enumerate(worksheet.columns, start=1):
            width = min(
                28,
                max(
                    10,
                    max(
                        len(str(cell.value)) if cell.value is not None else 0
                        for cell in cells
                    )
                    + 2,
                ),
            )
            worksheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(output_path)


def write_methodology_and_pipeline(
    summary: pd.DataFrame,
    filter_counts: pd.DataFrame,
    solver_counts: pd.DataFrame,
    output_path: Path,
) -> None:
    overall = summary[summary["ScopeType"] == "Overall"].iloc[0]
    lines = [
        "# Independent IV Robustness Check, 2018–2024",
        "",
        "## Methodology",
        "",
        "The primary paper measure remains OptionMetrics implied volatility. "
        "Self-calculated IV is used only as a validation exercise and is not "
        "substituted into the eSSVI/RND estimates.",
        "",
        "For every independently eligible quote, the market price is the bid–ask "
        "midpoint. Spot is the index close on the observation date; time to "
        "maturity is actual calendar days divided by 365; the risk-free rate is "
        "linearly interpolated from the OptionMetrics zero curve for the same "
        "date and maturity; and the dividend yield is the corresponding daily "
        "OptionMetrics index dividend yield.",
        "",
        "The calculation reuses the existing Black–Scholes call/put functions, "
        "analytical Vega, no-arbitrage bounds, multiple Newton starting values, "
        "step-halving, and Brent fallback in `part1_analysis.py`. No solver logic "
        "was rewritten.",
        "",
        "## Data flow",
        "",
        "```mermaid",
        "flowchart LR",
        "  A[Raw annual OptionMetrics quotes] --> B[Main-paper month-end dates and expirations]",
        "  B --> C[Quote filters independent of OM IV]",
        "  C --> D[Daily spot, zero rate, dividend yield, T and forward]",
        "  D --> E[OTM selection and duplicate resolution]",
        "  E --> F[Existing Newton/Brent IV solver]",
        "  F --> G[All self-IV attempts and failure diagnostics]",
        "  G --> H[Intersection: converged self IV, valid OM IV, main-sample delta]",
        "  H --> I[Bias, MAE, RMSE, correlation and figures]",
        "  E --> J[Primary OM IV branch]",
        "  J --> K[Existing eSSVI/RND pipeline, unchanged]",
        "```",
        "",
        "## Overall result",
        "",
        f"- Attempted independent IV inversions: {int(overall['NAttempted']):,}.",
        f"- Converged self-calculated IVs: {int(overall['NSelfIVConverged']):,} "
        f"({100*overall['ConvergenceRate']:.2f}%).",
        f"- Common valid comparison sample: {int(overall['NCommonValid']):,}.",
        f"- Mean error (self minus OptionMetrics): "
        f"{100*overall['MeanErrorBias']:.4f} volatility points.",
        f"- MAE: {100*overall['MAE']:.4f} points; "
        f"RMSE: {100*overall['RMSE']:.4f} points.",
        f"- Median absolute error: "
        f"{100*overall['MedianAbsoluteError']:.4f} points.",
        f"- Correlation: {overall['Correlation']:.6f}.",
        "",
        "## Methodological assumptions and limitations",
        "",
        "1. OptionMetrics IV may differ because its exact proprietary conventions "
        "for forward estimation, interest-rate interpolation, settlement timing, "
        "and quote handling are not fully replicated.",
        "2. The independent calculation uses the quoted bid–ask midpoint and the "
        "daily index close. Asynchronous option and index timestamps can create "
        "differences even when both calculations are correct.",
        "3. Zero rates and dividend yields are treated as continuously compounded "
        "annual rates, matching the Black–Scholes implementation used in the "
        "project.",
        "4. Actual calendar days/365 is used. AM settlement and holidays can "
        "produce small timing differences from OptionMetrics.",
        "5. The independent inversion sample does not require OptionMetrics IV "
        "or OptionMetrics delta. The paper comparison sample subsequently applies "
        "the main-sample 0.05–0.95 absolute-delta filter and intersects valid "
        "OptionMetrics IV with converged self IV. Independent solver failures "
        "remain in the robustness dataset and solver diagnostics.",
        "6. The self-IV exercise does not alter the existing OptionMetrics-IV "
        "eSSVI estimates or previously generated report.",
    ]
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def run(
    output_root: Path,
    *,
    refresh_extract: bool = False,
    refresh_solve: bool = False,
) -> None:
    output_root.mkdir(parents=True, exist_ok=True)
    classified = extract_independent_sample(
        output_root, force=refresh_extract
    )
    solved = calculate_independent_iv(
        classified, output_root, force=(refresh_solve or refresh_extract)
    )
    dataset = format_robustness_dataset(solved)
    summary, moneyness, yearly = build_summary_tables(dataset)
    (
        filter_counts,
        solver_counts,
        annual_filter,
        comparison_eligibility,
    ) = build_diagnostics(classified, dataset)
    codebook = diagnostic_codebook()

    csv_dir = output_root / "csv"
    chart_dir = output_root / "charts"
    csv_dir.mkdir(parents=True, exist_ok=True)
    dataset.to_csv(
        csv_dir / "iv_robustness_observations.csv.gz",
        index=False,
        compression="gzip",
        date_format="%Y-%m-%d",
    )
    dataset[
        dataset["Converged"]
        & dataset["OptionMetricsIVValid"]
        & dataset["MainSampleDeltaEligible"]
    ].to_csv(
        csv_dir / "common_valid_iv_comparison.csv.gz",
        index=False,
        compression="gzip",
        date_format="%Y-%m-%d",
    )
    classified[
        [
            "date",
            "Market",
            "exdate",
            "TargetTenorDays",
            "cp_flag",
            "Strike",
            "FilterStatus",
            "OptionMetricsIVValid",
            "MainSampleDeltaEligible",
        ]
    ].to_csv(
        csv_dir / "row_level_filter_audit.csv.gz",
        index=False,
        compression="gzip",
        date_format="%Y-%m-%d",
    )
    summary.to_csv(csv_dir / "iv_robustness_summary.csv", index=False)
    moneyness.to_csv(csv_dir / "moneyness_error_summary.csv", index=False)
    yearly.to_csv(csv_dir / "yearly_iv_error_summary.csv", index=False)
    filter_counts.to_csv(csv_dir / "filter_diagnostics.csv", index=False)
    solver_counts.to_csv(csv_dir / "solver_diagnostics.csv", index=False)
    annual_filter.to_csv(csv_dir / "annual_filter_diagnostics.csv", index=False)
    comparison_eligibility.to_csv(
        csv_dir / "comparison_eligibility_diagnostics.csv", index=False
    )
    codebook.to_csv(csv_dir / "diagnostic_codebook.csv", index=False)

    plot_iv_comparison(dataset, chart_dir)
    write_workbook(
        summary,
        moneyness,
        yearly,
        filter_counts,
        solver_counts,
        annual_filter,
        comparison_eligibility,
        codebook,
        dataset,
        output_root / "SPX_NDX_IV_Robustness_2018_2024.xlsx",
    )
    write_methodology_and_pipeline(
        summary,
        filter_counts,
        solver_counts,
        output_root / "methodology_pipeline_and_results.md",
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--data-root",
        type=Path,
        required=True,
        help=(
            "Directory containing OptionMetrics_SPX_NDX_<year> and "
            "OptionMetrics_SPX_NDX_auxiliary_2018_2024 folders."
        ),
    )
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--refresh-extract", action="store_true")
    parser.add_argument("--refresh-solve", action="store_true")
    args = parser.parse_args()
    configure_data_root(args.data_root)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=RuntimeWarning)
        run(
            args.output_root,
            refresh_extract=args.refresh_extract,
            refresh_solve=args.refresh_solve,
        )


if __name__ == "__main__":
    main()
