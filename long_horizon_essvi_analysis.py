#!/usr/bin/env python3
"""2018--2024 SPX/NDX eSSVI and risk-neutral-density analysis.

Design
------
* Monthly snapshots: last trading day of every month.
* Standard horizons: nearest listed expiry to 30, 90, 180, and 360 days.
* Quotes: OTM puts below the forward and OTM calls at/above the forward.
* Surface: eSSVI with sufficient butterfly constraints and sequential
  calendar-total-variance constraints.
* Outputs: fitted parameters, 25-delta risk reversal/butterfly, RND moments,
  annual summaries, explanatory regressions, charts, and an Excel workbook.
"""

from __future__ import annotations

import argparse
import math
import warnings
from pathlib import Path
from typing import Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy.optimize import brentq
from scipy.special import ndtr, ndtri

from essvi_rnd_analysis import (
    ESSVIParameters,
    butterfly_g,
    essvi_derivatives,
    essvi_total_variance,
    fit_slice,
)


SECID_TO_MARKET = {108105: "SPX", 102480: "NDX"}
YEARS = tuple(range(2018, 2025))
TARGET_TENORS = (30, 90, 180, 360)
TENOR_TOLERANCE = {30: 12, 90: 22, 180: 35, 360: 55}
SOURCE_ROOT: Path | None = None
AUXILIARY_DIR: Path | None = None
DEFAULT_OUTPUT = Path("outputs/long_horizon_essvi")

OPTION_COLUMNS = [
    "secid",
    "date",
    "exdate",
    "cp_flag",
    "strike_price",
    "best_bid",
    "best_offer",
    "volume",
    "open_interest",
    "impl_volatility",
    "delta",
]


def configure_data_root(source_root: Path) -> None:
    """Configure the directory containing the annual and auxiliary datasets."""
    global SOURCE_ROOT, AUXILIARY_DIR
    SOURCE_ROOT = source_root.expanduser().resolve()
    AUXILIARY_DIR = (
        SOURCE_ROOT / "OptionMetrics_SPX_NDX_auxiliary_2018_2024"
    )


def require_data_root() -> tuple[Path, Path]:
    if SOURCE_ROOT is None or AUXILIARY_DIR is None:
        raise RuntimeError(
            "OptionMetrics data root is not configured. "
            "Pass --data-root when running the script."
        )
    return SOURCE_ROOT, AUXILIARY_DIR


def option_path(year: int) -> Path:
    source_root, _ = require_data_root()
    return (
        source_root
        / f"OptionMetrics_SPX_NDX_{year}"
        / f"opprcd{year}.csv"
    )


def security_path(year: int) -> Path:
    source_root, _ = require_data_root()
    return (
        source_root
        / f"OptionMetrics_SPX_NDX_{year}"
        / f"secprd{year}.csv"
    )


def month_end_dates(security: pd.DataFrame) -> list[str]:
    dates = pd.to_datetime(security["date"])
    table = pd.DataFrame({"date": dates, "month": dates.dt.to_period("M")})
    return (
        table.groupby("month", sort=True)["date"]
        .max()
        .dt.strftime("%Y-%m-%d")
        .tolist()
    )


def load_auxiliary_data() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    _, auxiliary_dir = require_data_root()
    dividends = pd.read_csv(auxiliary_dir / "idxdvd.csv", parse_dates=["date"])
    zero_curve = pd.read_csv(auxiliary_dir / "zerocd.csv", parse_dates=["date"])
    vix = pd.read_csv(
        auxiliary_dir / "vix_2018_2024.csv", parse_dates=["date"]
    ).rename(columns={"close": "VIXClose"})
    dividends["DividendYield"] = dividends["rate"] / 100.0
    zero_curve["ZeroRate"] = zero_curve["rate"] / 100.0
    return dividends, zero_curve, vix[["date", "VIXClose"]]


def interpolate_zero_rate(
    zero_by_date: dict[pd.Timestamp, tuple[np.ndarray, np.ndarray]],
    date: pd.Timestamp,
    days: float,
) -> float:
    curve = zero_by_date.get(pd.Timestamp(date))
    if curve is None:
        return math.nan
    curve_days, rates = curve
    return float(np.interp(days, curve_days, rates))


def extract_month_end_quotes(
    output_root: Path,
    *,
    force: bool = False,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Stream large annual files and retain only standardized monthly slices."""
    cache = output_root / "cache" / "monthly_selected_quotes.pkl.gz"
    security_cache = output_root / "cache" / "daily_security.pkl.gz"
    audit_cache = output_root / "cache" / "extraction_audit.csv"
    if cache.exists() and security_cache.exists() and audit_cache.exists() and not force:
        return (
            pd.read_pickle(cache, compression="gzip"),
            pd.read_pickle(security_cache, compression="gzip"),
            pd.read_csv(audit_cache),
        )

    dividends, zero_curve, _ = load_auxiliary_data()
    dividend_lookup = (
        dividends.set_index(["secid", "date"])["DividendYield"].to_dict()
    )
    zero_by_date = {
        pd.Timestamp(date): (
            group["days"].to_numpy(dtype=float),
            group["ZeroRate"].to_numpy(dtype=float),
        )
        for date, group in zero_curve.sort_values(["date", "days"]).groupby("date")
    }

    selected_years: list[pd.DataFrame] = []
    securities: list[pd.DataFrame] = []
    audit_rows: list[dict[str, object]] = []

    for year in YEARS:
        source = option_path(year)
        security = pd.read_csv(security_path(year), parse_dates=["date"])
        security["Market"] = security["secid"].map(SECID_TO_MARKET)
        security = security[security["Market"].notna()].copy()
        securities.append(security)
        endpoint_strings = set(month_end_dates(security))
        retained_chunks: list[pd.DataFrame] = []
        rows_read = 0

        for chunk in pd.read_csv(
            source,
            usecols=OPTION_COLUMNS,
            chunksize=500_000,
            low_memory=False,
        ):
            rows_read += len(chunk)
            chunk = chunk[chunk["date"].isin(endpoint_strings)]
            if not chunk.empty:
                retained_chunks.append(chunk)

        if not retained_chunks:
            continue
        raw = pd.concat(retained_chunks, ignore_index=True)
        raw["date"] = pd.to_datetime(raw["date"])
        raw["exdate"] = pd.to_datetime(raw["exdate"])
        raw["DaysToMaturity"] = (raw["exdate"] - raw["date"]).dt.days
        numeric_columns = [
            "strike_price",
            "best_bid",
            "best_offer",
            "volume",
            "open_interest",
            "impl_volatility",
            "delta",
        ]
        for column in numeric_columns:
            raw[column] = pd.to_numeric(raw[column], errors="coerce")
        raw["Strike"] = raw["strike_price"] / 1000.0
        raw["MidPrice"] = 0.5 * (raw["best_bid"] + raw["best_offer"])
        raw["RelativeSpread"] = (
            (raw["best_offer"] - raw["best_bid"]) / raw["MidPrice"]
        )
        raw = raw[
            raw["secid"].isin(SECID_TO_MARKET)
            & raw["cp_flag"].isin(["C", "P"])
            & raw["DaysToMaturity"].between(7, 430)
            & (raw["best_bid"] >= 0)
            & (raw["best_offer"] > raw["best_bid"])
            & raw["MidPrice"].gt(0)
            & raw["RelativeSpread"].between(0, 0.50)
            & raw["impl_volatility"].between(0.01, 3.00)
            & raw["delta"].abs().between(0.05, 0.95)
            & raw["Strike"].gt(0)
        ].copy()

        expiration_choices: list[dict[str, object]] = []
        available = raw[
            ["date", "secid", "exdate", "DaysToMaturity"]
        ].drop_duplicates()
        for (date, secid), group in available.groupby(["date", "secid"]):
            for target in TARGET_TENORS:
                ordered = group.assign(
                    Distance=(group["DaysToMaturity"] - target).abs()
                ).sort_values(["Distance", "DaysToMaturity"])
                if ordered.empty:
                    continue
                choice = ordered.iloc[0]
                if float(choice["Distance"]) <= TENOR_TOLERANCE[target]:
                    expiration_choices.append(
                        {
                            "date": date,
                            "secid": secid,
                            "exdate": choice["exdate"],
                            "TargetTenorDays": target,
                        }
                    )
        choices = pd.DataFrame(expiration_choices)
        selected = raw.merge(
            choices, on=["date", "secid", "exdate"], how="inner"
        )
        spot_table = security.rename(
            columns={"close": "Spot"}
        )[["secid", "date", "Spot"]]
        selected = selected.merge(spot_table, on=["secid", "date"], how="left")
        selected["RiskFreeRate"] = [
            interpolate_zero_rate(zero_by_date, date, days)
            for date, days in zip(selected["date"], selected["DaysToMaturity"])
        ]
        selected["DividendYield"] = [
            dividend_lookup.get((int(secid), pd.Timestamp(date)), math.nan)
            for secid, date in zip(selected["secid"], selected["date"])
        ]
        selected["TimeToMaturityYears"] = selected["DaysToMaturity"] / 365.0
        selected["Forward"] = selected["Spot"] * np.exp(
            (selected["RiskFreeRate"] - selected["DividendYield"])
            * selected["TimeToMaturityYears"]
        )
        selected = selected[
            selected[
                ["Spot", "RiskFreeRate", "DividendYield", "Forward"]
            ].notna().all(axis=1)
        ].copy()
        selected["Market"] = selected["secid"].map(SECID_TO_MARKET)

        use_put = (selected["cp_flag"] == "P") & (
            selected["Strike"] < selected["Forward"]
        )
        use_call = (selected["cp_flag"] == "C") & (
            selected["Strike"] >= selected["Forward"]
        )
        selected = selected[use_put | use_call].copy()
        selected["OptionType"] = selected["cp_flag"].map(
            {"C": "Call", "P": "Put"}
        )
        selected["ImpliedVolatility"] = selected["impl_volatility"]
        selected["Maturity"] = selected["exdate"]
        selected["LogMoneyness"] = np.log(
            selected["Strike"] / selected["Forward"]
        )
        selected["ObservedTotalVariance"] = (
            selected["ImpliedVolatility"] ** 2
            * selected["TimeToMaturityYears"]
        )

        selected = selected.sort_values(
            [
                "date",
                "secid",
                "TargetTenorDays",
                "OptionType",
                "Strike",
                "RelativeSpread",
                "open_interest",
            ],
            ascending=[True, True, True, True, True, True, False],
        )
        selected = selected.drop_duplicates(
            [
                "date",
                "secid",
                "TargetTenorDays",
                "OptionType",
                "Strike",
            ],
            keep="first",
        )
        selected_years.append(selected)
        audit_rows.append(
            {
                "Year": year,
                "RawOptionRows": rows_read,
                "MonthEndRowsBeforeFilters": sum(
                    len(frame) for frame in retained_chunks
                ),
                "SelectedOTMRows": len(selected),
                "SnapshotDates": selected["date"].nunique(),
                "Slices": selected[
                    ["date", "Market", "TargetTenorDays"]
                ].drop_duplicates().shape[0],
            }
        )

    quotes = pd.concat(selected_years, ignore_index=True)
    daily_security = (
        pd.concat(securities, ignore_index=True)
        .sort_values(["secid", "date"])
        .reset_index(drop=True)
    )
    audit = pd.DataFrame(audit_rows)
    cache.parent.mkdir(parents=True, exist_ok=True)
    quotes.to_pickle(cache, compression="gzip")
    daily_security.to_pickle(security_cache, compression="gzip")
    audit.to_csv(audit_cache, index=False)
    return quotes, daily_security, audit


def density_summary(
    params: ESSVIParameters,
    maturity_years: float,
    forward: float,
) -> dict[str, float]:
    # Use a wide integration domain so long-dated/high-volatility slices retain
    # essentially all probability mass and preserve the forward first moment.
    normalized_strike = np.linspace(0.05, 3.00, 2951)
    strike = forward * normalized_strike
    k = np.log(normalized_strike)
    w, _, _ = essvi_derivatives(k, params.theta, params.rho, params.psi)
    g = butterfly_g(k, params.theta, params.rho, params.psi)
    d_minus = -k / np.sqrt(w) - 0.5 * np.sqrt(w)
    density = (
        np.exp(-0.5 * d_minus**2)
        / math.sqrt(2.0 * math.pi)
        * g
        / (strike * np.sqrt(w))
    )
    mass = float(np.trapezoid(density, strike))
    normalized = density / mass
    mean = float(np.trapezoid(strike * normalized, strike))
    variance = float(np.trapezoid((strike - mean) ** 2 * normalized, strike))
    standard_deviation = math.sqrt(max(variance, 0.0))
    skewness = float(
        np.trapezoid(
            ((strike - mean) / standard_deviation) ** 3 * normalized, strike
        )
    )
    excess_kurtosis = float(
        np.trapezoid(
            ((strike - mean) / standard_deviation) ** 4 * normalized, strike
        )
        - 3.0
    )
    low = normalized_strike <= 0.80
    high = normalized_strike >= 1.20
    return {
        "RNDMassOnGrid": mass,
        "RNDMeanOverForward": mean / forward,
        "RNDCoefficientOfVariation": standard_deviation / mean,
        "RNDSkewness": skewness,
        "RNDExcessKurtosis": excess_kurtosis,
        "RNDProbabilityBelow80Forward": float(
            np.trapezoid(normalized[low], strike[low])
        ),
        "RNDProbabilityAbove120Forward": float(
            np.trapezoid(normalized[high], strike[high])
        ),
        "RNDMinimumDensity": float(np.min(density)),
        "RNDMinimumG": float(np.min(g)),
    }


def delta_smile_metrics(
    params: ESSVIParameters, maturity_years: float
) -> dict[str, float]:
    def d1(k: float) -> float:
        w = float(
            essvi_total_variance(
                np.array([k]), params.theta, params.rho, params.psi
            )[0]
        )
        return -k / math.sqrt(w) + 0.5 * math.sqrt(w)

    def solve(target_d1: float) -> float:
        function = lambda value: d1(value) - target_d1
        grid = np.linspace(-2.0, 2.0, 801)
        values = np.array([function(value) for value in grid])
        crossings = np.where(values[:-1] * values[1:] <= 0)[0]
        if len(crossings) == 0:
            return math.nan
        index = int(crossings[np.argmin(np.abs(grid[crossings]))])
        return float(brentq(function, grid[index], grid[index + 1]))

    put_k = solve(float(ndtri(0.75)))
    call_k = solve(float(ndtri(0.25)))

    def volatility(k: float) -> float:
        if not math.isfinite(k):
            return math.nan
        w = float(
            essvi_total_variance(
                np.array([k]), params.theta, params.rho, params.psi
            )[0]
        )
        return math.sqrt(w / maturity_years)

    put_iv = volatility(put_k)
    call_iv = volatility(call_k)
    atm_iv = math.sqrt(params.theta / maturity_years)
    return {
        "Put25DeltaLogMoneyness": put_k,
        "Call25DeltaLogMoneyness": call_k,
        "Put25DeltaIV": put_iv,
        "Call25DeltaIV": call_iv,
        "RiskReversal25": call_iv - put_iv,
        "Butterfly25": 0.5 * (call_iv + put_iv) - atm_iv,
    }


def calibrate_all_slices(
    quotes: pd.DataFrame,
    output_root: Path,
    *,
    force: bool = False,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    cache = output_root / "cache" / "slice_metrics.pkl.gz"
    fitted_cache = output_root / "cache" / "fitted_points.pkl.gz"
    if cache.exists() and fitted_cache.exists() and not force:
        return (
            pd.read_pickle(cache, compression="gzip"),
            pd.read_pickle(fitted_cache, compression="gzip"),
        )

    metrics: list[dict[str, object]] = []
    fitted_frames: list[pd.DataFrame] = []
    grouped_dates = quotes.groupby(["date", "Market"], sort=True)
    total_groups = grouped_dates.ngroups

    for group_number, ((date, market), date_group) in enumerate(
        grouped_dates, start=1
    ):
        previous: ESSVIParameters | None = None
        for target, group in date_group.groupby("TargetTenorDays", sort=True):
            if len(group) < 8 or group["Strike"].nunique() < 8:
                continue
            group = group.copy()
            try:
                params, diagnostics = fit_slice(group, previous)
            except RuntimeError:
                continue
            maturity_years = float(group["TimeToMaturityYears"].iloc[0])
            forward = float(group["Forward"].iloc[0])
            k = group["LogMoneyness"].to_numpy(dtype=float)
            group["FittedTotalVariance"] = essvi_total_variance(
                k, params.theta, params.rho, params.psi
            )
            group["FittedIV"] = np.sqrt(
                group["FittedTotalVariance"] / maturity_years
            )
            group["IVResidual"] = (
                group["ImpliedVolatility"] - group["FittedIV"]
            )
            fitted_frames.append(group)
            metrics.append(
                {
                    "Date": pd.Timestamp(date),
                    "Year": pd.Timestamp(date).year,
                    "Market": market,
                    "TargetTenorDays": int(target),
                    "Maturity": pd.Timestamp(group["Maturity"].iloc[0]),
                    "ActualDaysToMaturity": int(
                        group["DaysToMaturity"].iloc[0]
                    ),
                    "TimeToMaturityYears": maturity_years,
                    "Spot": float(group["Spot"].iloc[0]),
                    "Forward": forward,
                    "RiskFreeRate": float(group["RiskFreeRate"].iloc[0]),
                    "DividendYield": float(group["DividendYield"].iloc[0]),
                    "Observations": len(group),
                    "Theta": params.theta,
                    "Rho": params.rho,
                    "Psi": params.psi,
                    "ATMVolatility": math.sqrt(params.theta / maturity_years),
                    "MedianRelativeSpread": float(
                        group["RelativeSpread"].median()
                    ),
                    "TotalVolume": float(group["volume"].fillna(0).sum()),
                    "TotalOpenInterest": float(
                        group["open_interest"].fillna(0).sum()
                    ),
                    **diagnostics,
                    **delta_smile_metrics(params, maturity_years),
                    **density_summary(params, maturity_years, forward),
                }
            )
            previous = params
        if group_number % 12 == 0 or group_number == total_groups:
            print(
                f"calibrated {group_number}/{total_groups} "
                f"date-market groups",
                flush=True,
            )

    metric_table = pd.DataFrame(metrics).sort_values(
        ["Date", "Market", "TargetTenorDays"]
    )
    fitted = pd.concat(fitted_frames, ignore_index=True)
    cache.parent.mkdir(parents=True, exist_ok=True)
    metric_table.to_pickle(cache, compression="gzip")
    fitted.to_pickle(fitted_cache, compression="gzip")
    return metric_table, fitted


def add_explanatory_variables(
    metrics: pd.DataFrame,
    daily_security: pd.DataFrame,
) -> pd.DataFrame:
    _, _, vix = load_auxiliary_data()
    security = daily_security.copy().sort_values(["secid", "date"])
    security["DailyLogReturn"] = security.groupby("secid")["close"].transform(
        lambda series: np.log(series).diff()
    )
    security["RealizedVol21"] = security.groupby("secid")[
        "DailyLogReturn"
    ].transform(lambda series: series.rolling(21, min_periods=15).std() * math.sqrt(252))
    endpoint = security[
        security["date"].isin(pd.to_datetime(metrics["Date"].unique()))
    ][["date", "Market", "RealizedVol21"]].copy()
    endpoint = endpoint.merge(vix, on="date", how="left")
    result = metrics.merge(
        endpoint,
        left_on=["Date", "Market"],
        right_on=["date", "Market"],
        how="left",
    ).drop(columns=["date"])
    result = result.sort_values(["Market", "TargetTenorDays", "Date"])
    result["MonthlyIndexReturn"] = result.groupby(
        ["Market", "TargetTenorDays"]
    )["Spot"].pct_change()
    return result


def robust_ols(
    data: pd.DataFrame,
    dependent: str,
    predictors: list[str],
    market: str,
) -> list[dict[str, object]]:
    usable = data[[dependent] + predictors].dropna().copy()
    if len(usable) <= len(predictors) + 3:
        return []
    x_raw = usable[predictors].to_numpy(dtype=float)
    means = x_raw.mean(axis=0)
    standard_deviations = x_raw.std(axis=0, ddof=0)
    standard_deviations[standard_deviations == 0] = 1.0
    x_standardized = (x_raw - means) / standard_deviations
    x = np.column_stack([np.ones(len(usable)), x_standardized])
    y = usable[dependent].to_numpy(dtype=float)
    inverse_xx = np.linalg.pinv(x.T @ x)
    beta = inverse_xx @ x.T @ y
    residual = y - x @ beta
    n, parameter_count = x.shape
    meat = x.T @ (x * residual[:, None] ** 2)
    covariance = (
        n / max(n - parameter_count, 1)
    ) * inverse_xx @ meat @ inverse_xx
    standard_error = np.sqrt(np.maximum(np.diag(covariance), 0.0))
    t_statistic = np.divide(
        beta,
        standard_error,
        out=np.full_like(beta, np.nan),
        where=standard_error > 0,
    )
    centered = y - y.mean()
    r_squared = 1.0 - float(residual @ residual) / float(centered @ centered)
    rows = []
    for name, coefficient, error, statistic in zip(
        ["Intercept"] + predictors, beta, standard_error, t_statistic
    ):
        rows.append(
            {
                "Market": market,
                "DependentVariable": dependent,
                "Predictor": name,
                "Coefficient": coefficient,
                "HC1RobustSE": error,
                "TStatistic": statistic,
                "Observations": n,
                "RSquared": r_squared,
                "TenorDays": 90,
                "PredictorsStandardized": name != "Intercept",
            }
        )
    return rows


def build_summary_tables(
    metrics: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    annual = (
        metrics.groupby(["Year", "Market", "TargetTenorDays"])
        .agg(
            Snapshots=("Date", "count"),
            ATMVolatilityMean=("ATMVolatility", "mean"),
            ATMVolatilityMedian=("ATMVolatility", "median"),
            RhoMean=("Rho", "mean"),
            RiskReversal25Mean=("RiskReversal25", "mean"),
            Butterfly25Mean=("Butterfly25", "mean"),
            RNDSkewnessMean=("RNDSkewness", "mean"),
            RNDDownside80Mean=("RNDProbabilityBelow80Forward", "mean"),
            RNDDispersionMean=("RNDCoefficientOfVariation", "mean"),
            IVRMSEMean=("RMSE_IV", "mean"),
            MedianRelativeSpread=("MedianRelativeSpread", "median"),
        )
        .reset_index()
    )
    comparison_columns = [
        "ATMVolatility",
        "Rho",
        "RiskReversal25",
        "Butterfly25",
        "RNDSkewness",
        "RNDProbabilityBelow80Forward",
        "RNDCoefficientOfVariation",
    ]
    wide = metrics.pivot_table(
        index=["Date", "TargetTenorDays"],
        columns="Market",
        values=comparison_columns,
    )
    comparison = wide.copy()
    for column in comparison_columns:
        if (column, "SPX") in comparison.columns and (
            column,
            "NDX",
        ) in comparison.columns:
            comparison[(f"NDXMinusSPX_{column}", "Difference")] = (
                comparison[(column, "NDX")] - comparison[(column, "SPX")]
            )
    comparison.columns = [
        f"{first}_{second}" for first, second in comparison.columns
    ]
    comparison = comparison.reset_index()

    predictors = [
        "VIXClose",
        "RealizedVol21",
        "MonthlyIndexReturn",
        "MedianRelativeSpread",
    ]
    regression_rows: list[dict[str, object]] = []
    ninety = metrics[metrics["TargetTenorDays"] == 90]
    for market in ("SPX", "NDX"):
        subset = ninety[ninety["Market"] == market]
        for dependent in (
            "ATMVolatility",
            "RiskReversal25",
            "RNDProbabilityBelow80Forward",
        ):
            regression_rows.extend(
                robust_ols(subset, dependent, predictors, market)
            )
    regressions = pd.DataFrame(regression_rows)
    return annual, comparison, regressions


def representative_density(
    row: pd.Series,
) -> pd.DataFrame:
    params = ESSVIParameters(
        float(row["Theta"]), float(row["Rho"]), float(row["Psi"])
    )
    normalized_strike = np.linspace(0.55, 1.45, 901)
    k = np.log(normalized_strike)
    w, _, _ = essvi_derivatives(k, params.theta, params.rho, params.psi)
    g = butterfly_g(k, params.theta, params.rho, params.psi)
    d_minus = -k / np.sqrt(w) - 0.5 * np.sqrt(w)
    standardized_density = (
        np.exp(-0.5 * d_minus**2)
        / math.sqrt(2.0 * math.pi)
        * g
        / (normalized_strike * np.sqrt(w))
    )
    return pd.DataFrame(
        {
            "StrikeOverForward": normalized_strike,
            "StandardizedDensity": standardized_density,
        }
    )


def plot_outputs(
    metrics: pd.DataFrame,
    annual: pd.DataFrame,
    fitted: pd.DataFrame,
    output_root: Path,
) -> None:
    chart_dir = output_root / "charts"
    chart_dir.mkdir(parents=True, exist_ok=True)
    colors = {"SPX": "#1f77b4", "NDX": "#ff7f0e"}
    ninety = metrics[metrics["TargetTenorDays"] == 90]

    def time_series_plot(
        column: str, ylabel: str, title: str, filename: str, scale: float = 1.0
    ) -> None:
        fig, ax = plt.subplots(figsize=(13, 6.5))
        for market, group in ninety.groupby("Market"):
            ax.plot(
                group["Date"],
                scale * group[column],
                label=market,
                color=colors[market],
                lw=1.4,
            )
        ax.set(title=title, xlabel="Month-end date", ylabel=ylabel)
        ax.xaxis.set_major_locator(mdates.YearLocator())
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
        ax.grid(alpha=0.25)
        ax.legend()
        fig.tight_layout()
        fig.savefig(chart_dir / filename, dpi=220, bbox_inches="tight")
        plt.close(fig)

    time_series_plot(
        "ATMVolatility",
        "ATM implied volatility (%)",
        "SPX versus NDX 90-day eSSVI ATM volatility",
        "monthly_atm_volatility_90d.png",
        100.0,
    )
    time_series_plot(
        "RiskReversal25",
        "25-delta risk reversal (volatility points)",
        "SPX versus NDX 90-day 25-delta risk reversal",
        "monthly_risk_reversal_90d.png",
        100.0,
    )
    time_series_plot(
        "RNDProbabilityBelow80Forward",
        "Risk-neutral probability (%)",
        "Probability of terminal index below 80% of forward (90-day RND)",
        "monthly_rnd_downside_probability_90d.png",
        100.0,
    )

    fig, ax = plt.subplots(figsize=(12, 6.5))
    for market, group in annual.groupby("Market"):
        yearly = (
            group.groupby("Year")["ATMVolatilityMean"].mean().reset_index()
        )
        ax.plot(
            yearly["Year"],
            100 * yearly["ATMVolatilityMean"],
            marker="o",
            lw=2,
            label=market,
            color=colors[market],
        )
    ax.set(
        title="Annual mean eSSVI ATM volatility across standard tenors",
        xlabel="Year",
        ylabel="ATM implied volatility (%)",
    )
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(
        chart_dir / "annual_atm_volatility.png", dpi=220, bbox_inches="tight"
    )
    plt.close(fig)

    representative_targets = [
        pd.Timestamp("2020-03-31"),
        pd.Timestamp("2022-06-30"),
        pd.Timestamp("2024-12-31"),
    ]
    fig, axes = plt.subplots(1, 3, figsize=(16, 5), sharey=True)
    for ax, requested in zip(axes, representative_targets):
        available = ninety["Date"].drop_duplicates().sort_values()
        date = min(available, key=lambda value: abs(value - requested))
        for market in ("SPX", "NDX"):
            row = ninety[
                (ninety["Date"] == date) & (ninety["Market"] == market)
            ]
            if row.empty:
                continue
            density = representative_density(row.iloc[0])
            ax.plot(
                density["StrikeOverForward"],
                density["StandardizedDensity"],
                label=market,
                color=colors[market],
                lw=2,
            )
        ax.axvline(1.0, color="black", ls="--", lw=0.8)
        ax.set_title(pd.Timestamp(date).strftime("%Y-%m-%d"))
        ax.set_xlabel("Terminal level / forward")
        ax.grid(alpha=0.25)
    axes[0].set_ylabel("Risk-neutral density")
    axes[0].legend()
    fig.suptitle("Representative 90-day eSSVI risk-neutral densities")
    fig.tight_layout(rect=(0, 0, 1, 0.94))
    fig.savefig(
        chart_dir / "representative_rnd_comparison.png",
        dpi=220,
        bbox_inches="tight",
    )
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(14, 5.5), sharey=True)
    for ax, market in zip(axes, ("SPX", "NDX")):
        subset = ninety[ninety["Market"] == market]
        scatter = ax.scatter(
            subset["VIXClose"],
            100 * subset["RNDProbabilityBelow80Forward"],
            c=subset["Date"].dt.year,
            cmap="viridis",
            s=30,
            alpha=0.8,
        )
        ax.set_title(market)
        ax.set_xlabel("VIX close")
        ax.grid(alpha=0.25)
    axes[0].set_ylabel("RND probability below 80% forward (%)")
    fig.colorbar(scatter, ax=axes, label="Year")
    fig.suptitle("VIX and model-implied downside tail probability")
    fig.savefig(
        chart_dir / "vix_vs_rnd_downside.png",
        dpi=220,
        bbox_inches="tight",
    )
    plt.close(fig)

    dates = [
        min(
            fitted["date"].drop_duplicates(),
            key=lambda value: abs(value - target),
        )
        for target in representative_targets
    ]
    fig, axes = plt.subplots(2, 3, figsize=(16, 9), sharey="row")
    for column, date in enumerate(dates):
        for row_number, market in enumerate(("SPX", "NDX")):
            ax = axes[row_number, column]
            group = fitted[
                (fitted["date"] == date)
                & (fitted["Market"] == market)
                & (fitted["TargetTenorDays"] == 90)
            ].sort_values("LogMoneyness")
            if group.empty:
                continue
            ax.scatter(
                group["LogMoneyness"],
                100 * group["ImpliedVolatility"],
                s=15,
                alpha=0.65,
                label="Observed",
            )
            ax.plot(
                group["LogMoneyness"],
                100 * group["FittedIV"],
                color="black",
                lw=2,
                label="eSSVI",
            )
            ax.set_title(f"{market} {pd.Timestamp(date):%Y-%m-%d}")
            ax.set_xlabel("Log-forward moneyness")
            ax.grid(alpha=0.25)
    axes[0, 0].set_ylabel("IV (%)")
    axes[1, 0].set_ylabel("IV (%)")
    axes[0, 0].legend()
    fig.suptitle("Representative observed smiles and eSSVI fits (90-day target)")
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    fig.savefig(
        chart_dir / "representative_essvi_fits.png",
        dpi=220,
        bbox_inches="tight",
    )
    plt.close(fig)


def write_analysis_summary(
    metrics: pd.DataFrame,
    annual: pd.DataFrame,
    comparison: pd.DataFrame,
    regressions: pd.DataFrame,
    audit: pd.DataFrame,
    output_path: Path,
) -> None:
    ninety = metrics[metrics["TargetTenorDays"] == 90]
    by_market = ninety.groupby("Market")
    covid = ninety[
        ninety["Date"].between("2020-02-01", "2020-06-30")
    ].groupby("Market")
    ndx_minus_spx_atm = comparison[
        "NDXMinusSPX_ATMVolatility_Difference"
    ].dropna()
    ndx_minus_spx_downside = comparison[
        "NDXMinusSPX_RNDProbabilityBelow80Forward_Difference"
    ].dropna()

    def regression_line(market: str, dependent: str, predictor: str) -> str:
        row = regressions[
            (regressions["Market"] == market)
            & (regressions["DependentVariable"] == dependent)
            & (regressions["Predictor"] == predictor)
        ].iloc[0]
        return (
            f"{market}: β={row['Coefficient']:.4f}, "
            f"t={row['TStatistic']:.2f}"
        )

    lines = [
        "# SPX and NDX eSSVI/RND Results, 2018–2024",
        "",
        "## Sample and estimation",
        "",
        (
            f"- Processed {int(audit['RawOptionRows'].sum()):,} raw option rows "
            f"from seven complete calendar years."
        ),
        (
            f"- Retained {len(metrics)} successfully calibrated monthly "
            "market-tenor slices at 30, 90, 180, and 360-day target horizons."
        ),
        (
            f"- Median IV RMSE is {100*metrics['RMSE_IV'].median():.3f} "
            "volatility points; every retained slice has a non-negative "
            f"density factor (minimum g={metrics['MinDensityFactorG'].min():.6f})."
        ),
        "",
        "## Main cross-market findings",
        "",
        (
            f"- At the 90-day horizon, mean ATM IV is "
            f"{100*by_market['ATMVolatility'].mean()['SPX']:.2f}% for SPX and "
            f"{100*by_market['ATMVolatility'].mean()['NDX']:.2f}% for NDX."
        ),
        (
            f"- Across matched date-tenor observations, NDX ATM IV exceeds SPX "
            f"by {100*ndx_minus_spx_atm.mean():.2f} volatility points on average."
        ),
        (
            f"- NDX's probability of ending below 80% of forward exceeds SPX "
            f"by {100*ndx_minus_spx_downside.mean():.2f} percentage points on "
            "average across matched observations."
        ),
        (
            f"- Mean 90-day 25-delta risk reversal is "
            f"{100*by_market['RiskReversal25'].mean()['SPX']:.2f} points for SPX "
            f"and {100*by_market['RiskReversal25'].mean()['NDX']:.2f} points "
            "for NDX; more negative values indicate stronger downside skew."
        ),
        "",
        "## Stress-period evidence",
        "",
        (
            f"- During February–June 2020, mean 90-day ATM IV reaches "
            f"{100*covid['ATMVolatility'].mean()['SPX']:.2f}% for SPX and "
            f"{100*covid['ATMVolatility'].mean()['NDX']:.2f}% for NDX."
        ),
        (
            f"- Over the same period, the model-implied probability below 80% "
            f"of forward averages "
            f"{100*covid['RNDProbabilityBelow80Forward'].mean()['SPX']:.2f}% "
            f"for SPX and "
            f"{100*covid['RNDProbabilityBelow80Forward'].mean()['NDX']:.2f}% "
            "for NDX."
        ),
        "",
        "## Explanatory regressions (90-day horizon)",
        "",
        (
            "- Standardized VIX coefficient in the ATM-IV regression — "
            + regression_line("SPX", "ATMVolatility", "VIXClose")
            + "; "
            + regression_line("NDX", "ATMVolatility", "VIXClose")
            + "."
        ),
        (
            "- Standardized VIX coefficient in the downside-probability "
            "regression — "
            + regression_line(
                "SPX", "RNDProbabilityBelow80Forward", "VIXClose"
            )
            + "; "
            + regression_line(
                "NDX", "RNDProbabilityBelow80Forward", "VIXClose"
            )
            + "."
        ),
        "",
        "## Interpretation limits",
        "",
        "- RND tail probabilities are model-implied. They are more sensitive "
        "to eSSVI wing extrapolation than ATM volatility or risk reversal.",
        "- The regressions establish conditional associations, not causal "
        "effects. Sector composition and investor-type data would be needed "
        "for a complete causal decomposition.",
    ]
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_workbook(
    metrics: pd.DataFrame,
    annual: pd.DataFrame,
    comparison: pd.DataFrame,
    regressions: pd.DataFrame,
    audit: pd.DataFrame,
    output_path: Path,
) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        metrics.to_excel(writer, sheet_name="Monthly eSSVI RND", index=False)
        annual.to_excel(writer, sheet_name="Annual Summary", index=False)
        comparison.to_excel(writer, sheet_name="SPX NDX Comparison", index=False)
        regressions.to_excel(writer, sheet_name="Regressions", index=False)
        audit.to_excel(writer, sheet_name="Data Audit", index=False)
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(output_path)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center")
        for index, cells in enumerate(worksheet.columns, start=1):
            width = min(
                30,
                max(
                    10,
                    max(
                        len(str(cell.value)) if cell.value is not None else 0
                        for cell in cells
                    )
                    + 2,
                ),
            )
            worksheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(output_path)


def run(
    output_root: Path,
    *,
    refresh_extract: bool = False,
    refresh_fit: bool = False,
) -> None:
    output_root.mkdir(parents=True, exist_ok=True)
    quotes, daily_security, audit = extract_month_end_quotes(
        output_root, force=refresh_extract
    )
    metrics, fitted = calibrate_all_slices(
        quotes, output_root, force=(refresh_fit or refresh_extract)
    )
    metrics = add_explanatory_variables(metrics, daily_security)
    annual, comparison, regressions = build_summary_tables(metrics)

    csv_dir = output_root / "csv"
    csv_dir.mkdir(parents=True, exist_ok=True)
    metrics.to_csv(
        csv_dir / "monthly_essvi_rnd_metrics.csv",
        index=False,
        date_format="%Y-%m-%d",
    )
    fitted.to_csv(
        csv_dir / "essvi_fitted_points.csv",
        index=False,
        date_format="%Y-%m-%d",
    )
    annual.to_csv(csv_dir / "annual_summary.csv", index=False)
    comparison.to_csv(
        csv_dir / "spx_ndx_matched_comparison.csv",
        index=False,
        date_format="%Y-%m-%d",
    )
    regressions.to_csv(csv_dir / "explanatory_regressions.csv", index=False)
    audit.to_csv(csv_dir / "data_audit.csv", index=False)

    plot_outputs(metrics, annual, fitted, output_root)
    write_workbook(
        metrics,
        annual,
        comparison,
        regressions,
        audit,
        output_root / "SPX_NDX_eSSVI_RND_2018_2024.xlsx",
    )
    write_analysis_summary(
        metrics,
        annual,
        comparison,
        regressions,
        audit,
        output_root / "paper_ready_results.md",
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--data-root",
        type=Path,
        required=True,
        help=(
            "Directory containing OptionMetrics_SPX_NDX_<year> and "
            "OptionMetrics_SPX_NDX_auxiliary_2018_2024 folders."
        ),
    )
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--refresh-extract", action="store_true")
    parser.add_argument("--refresh-fit", action="store_true")
    args = parser.parse_args()
    configure_data_root(args.data_root)
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Values in x were outside bounds during a minimize step",
        )
        run(
            args.output_root,
            refresh_extract=args.refresh_extract,
            refresh_fit=args.refresh_fit,
        )


if __name__ == "__main__":
    main()
